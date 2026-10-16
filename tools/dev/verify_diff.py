# -*- coding: utf-8 -*-
"""验证差异表重排"""
from openpyxl import load_workbook
p = r'D:\openLM\差异对比_已更新.xlsx'
wb = load_workbook(p)
ws = wb['功能对比']
for r in range(3, ws.max_row + 1):
    v1 = ws.cell(r, 2).value
    v6 = ws.cell(r, 6).value
    f = ws.cell(r, 6).fill.fgColor.rgb
    if v1:
        print(r, '|', str(v1)[:34], '|', v6, '|', f)
