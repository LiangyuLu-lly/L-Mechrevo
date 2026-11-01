# -*- coding: utf-8 -*-
"""读差异表当前功能对比 sheet 状态"""
from openpyxl import load_workbook
p = r'D:\openLM\差异对比_已更新.xlsx'
wb = load_workbook(p)
print('sheets:', wb.sheetnames)
ws = wb['功能对比']
print('dims:', ws.dimensions, 'max_row:', ws.max_row, 'max_col:', ws.max_column)
# 表头
for r in range(1, 3):
    print('row', r, [ws.cell(r, c).value for c in range(1, ws.max_column+1)])
# 数据行：序号 功能 状态 列
for r in range(3, ws.max_row+1):
    vals = [ws.cell(r, c).value for c in range(1, min(ws.max_column+1, 6))]
    if any(v is not None for v in vals):
        fills = []
        for c in range(1, ws.max_column+1):
            f = ws.cell(r, c).fill
            fills.append(f.fgColor.rgb if f and f.fgColor and f.fgColor.rgb not in (None, '00000000') else '')
        print(r, vals, fills)
