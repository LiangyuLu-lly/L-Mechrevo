# -*- coding: utf-8 -*-
"""自定义模式编辑器 4 项 -> 已加入 + 重排 + 状态标色"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

XLSX = r'C:\Users\xiaohan\Desktop\L-Mechrevo-作者beta21-vs-群友版0.289.6-差异对比.xlsx'
wb = openpyxl.load_workbook(XLSX)
old = wb['功能对比']

# 1) 更新 4 项（按功能项关键词定位）
targets = {
    '托盘列全部模式': ('有', '有（0.290.0 合并）', '已完整合并（CustomModeForm + PerfModeCollection 12 档）', '已加入'),
    '内置模式二级自定义': ('有（beta20 引入，beta21 被编辑器取代）', '有（0.290.0 合并）', '已合并（电源计划/电源模式/睿频/风扇增强/刷新率/OC）', '已加入'),
    '自定义档命名写回确认': ('有', '有（0.290.0 合并，≤12 字符）', '已合并（命名限长已对齐 12 字符）', '已加入'),
    '自定义模式切后': ('有', '有（0.290.0 合并）', '已合并（PerfApplyVerify 逐项回读校验）', '已加入'),
}
updated = []
for r in range(3, old.max_row + 1):
    b = str(old.cell(r, 2).value or '')
    for kw, vals in targets.items():
        if kw in b:
            for c, v in enumerate(vals, start=3):
                old.cell(r, c).value = v
            updated.append((r, kw))
            break
print('OK 更新 4 项:', updated)

# 2) 读取数据行 -> 分组 -> 重建
rows = []
for r in range(3, old.max_row + 1):
    vals = [old.cell(r, c).value for c in range(1, 7)]
    if all(v in (None, '') for v in vals):
        continue
    rows.append((str(vals[5] or ''), vals))

groups = {'未加入': [], '有差异': [], '已加入': []}
for status, vals in rows:
    key = status if status in groups else '已加入'
    groups[key].append(vals)
order = groups['未加入'] + groups['有差异'] + groups['已加入']

del wb['功能对比']
ws = wb.create_sheet('功能对比', 1)

ws.cell(1, 1, '功能对比：作者 beta21  vs  群友版 0.290.0 —— 未加入在上 · 有差异居中 · 已加入在底（绿）')
ws.merge_cells('A1:F1')
ws.cell(1, 1).font = Font(bold=True, size=12)
ws.cell(1, 1).alignment = Alignment(horizontal='center', vertical='center')
ws.row_dimensions[1].height = 24

headers = ['模块', '功能项', '作者 beta21', '群友版', '说明 / 备注', '已加入']
for c, h in enumerate(headers, start=1):
    cell = ws.cell(2, c, h)
    cell.font = Font(bold=True)
    cell.alignment = Alignment(horizontal='center', vertical='center')
    cell.fill = PatternFill('solid', fgColor='D9E2F3')
ws.row_dimensions[2].height = 22

STATUS_FILL = {
    '未加入': PatternFill('solid', fgColor='FDE9E9'),
    '有差异': PatternFill('solid', fgColor='FDF3DC'),
    '已加入': PatternFill('solid', fgColor='E9F3EC'),
}
for i, vals in enumerate(order):
    r = 3 + i
    st = str(vals[5] or '')
    for c, v in enumerate(vals, start=1):
        cell = ws.cell(r, c, v)
        if c == 6 and st in STATUS_FILL:
            cell.fill = STATUS_FILL[st]
        cell.alignment = Alignment(vertical='center', wrap_text=(c == 2))

widths = [12, 46, 20, 30, 34, 10]
for c, w in enumerate(widths, start=1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = w

print('未加入=%d 有差异=%d 已加入=%d' % (
    len(groups['未加入']), len(groups['有差异']), len(groups['已加入'])))

try:
    wb.save(XLSX)
    print('SAVED OK')
except PermissionError:
    print('LOCKED - save failed')
