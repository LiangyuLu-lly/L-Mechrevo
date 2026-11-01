# -*- coding: utf-8 -*-
"""更新差异表格：功能对比在合并区后追加、独有功能汇总追加"""
import openpyxl

XLSX = r'C:\Users\xiaohan\Desktop\L-Mechrevo-作者beta21-vs-群友版0.289.6-差异对比.xlsx'
wb = openpyxl.load_workbook(XLSX)

ws1 = wb['功能对比']
new_rows = [
    ('电源与系统', '睡眠省电 6 开关（注册表直写当前电源方案）', '无', '有（0.290.0）', '群友独有（OpenRevo 抓包实现）', '已加入'),
    ('电源与系统', '插电/离电（AC/DC）分档切换', '无', '有（0.290.0）', '群友独有', '已加入'),
    ('电源与系统', '睡眠开关悬停解释提示（ToolTip）', '无', '有（0.290.0）', '群友独有', '已加入'),
    ('电源与系统', '睡眠关闭值对齐 Windows 语义（由Windows管理/仅重要唤醒源）', '无', '有（0.290.0）', '群友独有', '已加入'),
]

# 直接追加在现有最大行之后（52-55 是合并空行，跳过）
start = ws1.max_row + 1
for i, row in enumerate(new_rows):
    for c, v in enumerate(row, start=1):
        ws1.cell(start + i, c).value = v
print('OK 功能对比追加', start, '-', start + len(new_rows) - 1)

ws2 = wb['独有功能汇总']
r2 = ws2.max_row
while r2 > 1 and (ws2.cell(r2, 1).value in (None, '') and ws2.cell(r2, 2).value in (None, '')):
    r2 -= 1
ws2.cell(r2 + 1, 2).value = '睡眠省电 6 开关（AC/DC 分档 + 悬停解释，0.290.0）'
print('OK 独有功能汇总追加')

wb.save(XLSX)
print('ALL DONE')
