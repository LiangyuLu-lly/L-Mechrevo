# -*- coding: utf-8 -*-
"""功能对比：仅修正过时条目（11/12），不删行，输出到临时文件"""
import openpyxl

XLSX = r'C:\Users\xiaohan\Desktop\L-Mechrevo-作者beta21-vs-群友版0.289.6-差异对比.xlsx'
TMP = r'D:\openLM\差异对比_已更新.xlsx'

wb = openpyxl.load_workbook(XLSX)
ws1 = wb['功能对比']

fixes = {
    11: ('电池与充电',
         '充电上限：观察实际充电行为确认生效，不受控自动恢复满充并隐藏',
         '有',
         '有（18.6：40系三档、50系观察+回退三档不隐藏）',
         '实现不同：作者观察+不受控隐藏，群友 50 系观察+回退三档',
         '有差异'),
    12: ('托盘与界面',
         '温度/风扇读数过期显示"—"，停转显示 0%',
         '有',
         '有（18.6 移植）',
         '已移植实现',
         '已加入'),
}
for r, row in fixes.items():
    for c, v in enumerate(row, start=1):
        ws1.cell(r, c).value = v
print('OK 修正行 11/12')

wb.save(TMP)
print('saved tmp OK')

# 验证
wb2 = openpyxl.load_workbook(TMP)
ws = wb2['功能对比']
print('max_row:', ws.max_row)
for r in list(range(10, 16)) + list(range(50, 60)):
    a = ws.cell(r, 1).value or ''
    b = ws.cell(r, 2).value or ''
    f = ws.cell(r, 6).value or ''
    print(r, '|', a[:10], '||', b[:36], '||', f)
print('DONE')
